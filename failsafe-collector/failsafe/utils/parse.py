"""Pre-built parse utility functions.

Deterministic extraction logic for HTML, JSON, PDF, and Excel sources.
"""

from __future__ import annotations

import io
import re
from typing import Any, Dict, List, Optional

from bs4 import BeautifulSoup, Tag


def select_elements(soup: BeautifulSoup | Tag, selector: str) -> list[Tag]:
    return soup.select(selector)


def select_one(soup: BeautifulSoup | Tag, selector: str) -> Optional[Tag]:
    return soup.select_one(selector)


def extract_text(element: Optional[Tag], strip: bool = True) -> str:
    if element is None:
        return ""
    text = element.get_text(separator=" ", strip=strip)
    return text


def extract_attr(element: Optional[Tag], attr: str) -> str:
    if element is None:
        return ""
    return element.get(attr, "") or ""


def extract_href(element: Optional[Tag], base_url: str = "") -> str:
    href = extract_attr(element, "href")
    if href and not href.startswith(("http://", "https://")):
        href = base_url.rstrip("/") + "/" + href.lstrip("/")
    return href


def extract_fields_from_page(
    soup: BeautifulSoup | Tag,
    field_selectors: dict[str, str],
    base_url: str = "",
) -> dict[str, str]:
    result = {}
    for field_name, selector in field_selectors.items():
        if selector.startswith("@"):
            attr_name = selector.split("@", 1)[1]
            parts = attr_name.split("|", 1)
            if len(parts) == 2:
                sel, attr = parts
                el = select_one(soup, sel)
                result[field_name] = extract_attr(el, attr)
            else:
                result[field_name] = ""
        elif selector.startswith("href:"):
            sel = selector[5:]
            el = select_one(soup, sel)
            result[field_name] = extract_href(el, base_url)
        else:
            el = select_one(soup, selector)
            result[field_name] = extract_text(el)
    return result


def extract_json_path(data: Any, path: str) -> Any:
    """Simple JSON path extraction (dot notation with array indexing).

    Examples: "$.data.items", "$.results[0].name", "$"
    """
    if path == "$":
        return data
    parts = path.lstrip("$").lstrip(".").split(".")
    current = data
    for part in parts:
        if not part:
            continue
        array_match = re.match(r"(\w+)\[(\d+)\]", part)
        if array_match:
            key, idx = array_match.group(1), int(array_match.group(2))
            if isinstance(current, dict):
                current = current.get(key, [])
            if isinstance(current, list) and idx < len(current):
                current = current[idx]
            else:
                return None
        elif isinstance(current, dict):
            current = current.get(part)
        elif isinstance(current, list):
            try:
                current = current[int(part)]
            except (ValueError, IndexError):
                return None
        else:
            return None
        if current is None:
            return None
    return current


def map_fields(record: dict[str, Any], mappings: dict[str, str]) -> dict[str, Any]:
    """Map source field names to target field names."""
    if not mappings:
        return record
    result = {}
    for target_name, source_path in mappings.items():
        value = extract_json_path(record, f"$.{source_path}") if "." in source_path else record.get(source_path)
        result[target_name] = value
    return result


def parse_pdf(content: bytes) -> str:
    from PyPDF2 import PdfReader

    reader = PdfReader(io.BytesIO(content))
    text_parts = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            text_parts.append(text)
    return "\n".join(text_parts)


def parse_excel(
    content: bytes,
    sheet_name: Optional[str] = None,
    start_row: int = 0,
    header_row: Optional[int] = 0,
) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    if sheet_name:
        ws = wb[sheet_name]
    else:
        ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    if header_row is not None:
        headers = [str(h) if h else f"col_{i}" for i, h in enumerate(rows[header_row])]
        data_start = max(header_row + 1, start_row)
    else:
        headers = [f"col_{i}" for i in range(len(rows[0]))]
        data_start = start_row

    records = []
    for row in rows[data_start:]:
        record = {}
        for i, val in enumerate(row):
            if i < len(headers):
                record[headers[i]] = val
        records.append(record)
    return records


def parse_csv_text(text: str, delimiter: str = ",") -> list[dict[str, Any]]:
    import csv

    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    return [dict(row) for row in reader]
